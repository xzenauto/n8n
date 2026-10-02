#!/usr/bin/env python3
"""Convierte directorio_cultura_municipal.csv en un Excel con pestañas:
Prioritarios (municipales con titular o contacto), Todos, Festivales y Resumen.
Requiere: pip install openpyxl
"""
import csv
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

entrada = sys.argv[1] if len(sys.argv) > 1 else "directorio_cultura_municipal.csv"
salida = sys.argv[2] if len(sys.argv) > 2 else "directorio_cultura_municipal.xlsx"

with open(entrada, encoding="utf-8-sig") as f:
    filas = list(csv.DictReader(f))
columnas = list(filas[0].keys())

ANCHOS = {"estado": 16, "municipio": 20, "tipo": 24, "nivel": 12, "institucion": 40,
          "adscripcion": 30, "titular": 30, "cargo_titular": 22, "telefono": 22,
          "correo": 38, "correo_institucional": 10, "pagina_web": 30, "domicilio": 40,
          "organizador_festival": 30, "fecha_festival": 26, "ficha_sic": 20,
          "ultima_actualizacion": 12}
ENCABEZADOS = {"correo_institucional": "correo gob/institucional",
               "ultima_actualizacion": "actualizado"}


def hoja(wb, titulo, datos, cols=columnas):
    ws = wb.create_sheet(titulo)
    ws.append([ENCABEZADOS.get(c, c).replace("_", " ").upper() for c in cols])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="7A1F3D")
    for d in datos:
        ws.append([d.get(c, "") for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = ANCHOS.get(c, 18)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook()
wb.remove(wb.active)

municipales = [f for f in filas if f["nivel"] == "Municipal" and f["tipo"] != "Festival / Feria cultural"]
prioritarios = [f for f in municipales if f["titular"] or f["correo"] or f["telefono"]]
prioritarios.sort(key=lambda f: (f["estado"], f["municipio"], f["tipo"] != "Instituto/Dirección municipal de cultura"))
festivales = [f for f in filas if f["tipo"] == "Festival / Feria cultural"]

hoja(wb, "Municipales (prioritarios)", prioritarios,
     [c for c in columnas if c not in ("organizador_festival", "fecha_festival", "adscripcion")])
hoja(wb, "Festivales y ferias", festivales,
     ["estado", "municipio", "institucion", "organizador_festival", "fecha_festival",
      "titular", "telefono", "correo", "correo_institucional", "pagina_web", "ficha_sic",
      "ultima_actualizacion"])
hoja(wb, "Todos los registros", filas)

res = defaultdict(lambda: defaultdict(int))
for f in filas:
    e = res[f["estado"] or "(sin estado)"]
    e["registros"] += 1
    e["municipales"] += f["nivel"] == "Municipal"
    e["con_titular"] += bool(f["titular"])
    e["con_correo"] += bool(f["correo"])
    e["con_telefono"] += bool(f["telefono"])
mun = defaultdict(set)
for f in municipales:
    mun[f["estado"] or "(sin estado)"].add(f["municipio"])
cols = ["estado", "registros", "municipales", "municipios_distintos", "con_titular", "con_correo", "con_telefono"]
resumen = [dict(estado=k, municipios_distintos=len(mun[k]), **v) for k, v in sorted(res.items())]
hoja(wb, "Resumen por estado", resumen, cols)

wb.save(salida)
print(f"{salida}: {len(prioritarios)} prioritarios, {len(festivales)} festivales, {len(filas)} total")
